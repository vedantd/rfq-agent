#!/usr/bin/env python3
"""
Generate the RudCo RFQ-to-Quote demo dataset.

Everything here is SYNTHETIC: supplier, mills, prices and market indices are
invented for a hackathon demo. Do not treat any number as real market data.

Run:  python3 tools/generate.py      (from the rudco_demo_data folder)
Re-running is deterministic (fixed random seed).
"""
import csv
import json
import math
import random
import sqlite3
from email.message import EmailMessage
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
random.seed(42)

AS_OF = "2026-10-05"
SUPPLIER = "Gowanus Metal Supply (fictional demo supplier)"
SYNTH = "SYNTHETIC demo data - not real market prices"

# --------------------------------------------------------------------------
# 1. Reference tables
# --------------------------------------------------------------------------
DENSITY = {  # lb / in^3
    "stainless steel": 0.289,
    "aluminum": 0.0975,
    "carbon steel": 0.2836,
    "galvanized steel": 0.2836,
}

# gauge -> (stainless, galvanized, carbon steel MSG, aluminum B&S) in inches
GAUGES = {
    10: (0.135, 0.1382, 0.1345, 0.1019),
    11: (0.120, 0.1233, 0.1196, 0.0907),
    12: (0.105, 0.1084, 0.1046, 0.0808),
    14: (0.075, 0.0785, 0.0747, 0.0641),
    16: (0.060, 0.0635, 0.0598, 0.0508),
    18: (0.048, 0.0516, 0.0478, 0.0403),
    20: (0.036, 0.0396, 0.0359, 0.0320),
    22: (0.030, 0.0336, 0.0299, 0.0253),
    24: (0.024, 0.0276, 0.0239, 0.0201),
}
TUBE_GAUGE_WALL = {11: 0.120, 13: 0.095, 14: 0.083, 16: 0.065, 18: 0.049}

# id, family, grade, form, finish_class, low, median, high  ($/lb)
BENCH = [
    ("BM-SS304-SH-2B", "stainless steel", "304", "sheet", "2B", 2.85, 3.20, 3.65),
    ("BM-SS304-SH-4", "stainless steel", "304", "sheet", "#4", 3.20, 3.58, 4.10),
    ("BM-SS304-SH-BA", "stainless steel", "304", "sheet", "BA", 3.10, 3.45, 3.95),
    ("BM-SS316-SH-2B", "stainless steel", "316", "sheet", "2B", 4.05, 4.50, 5.10),
    ("BM-SS316-SH-4", "stainless steel", "316", "sheet", "#4", 4.45, 4.90, 5.55),
    ("BM-SS304-PL", "stainless steel", "304", "plate", "No.1 HRAP", 2.95, 3.30, 3.80),
    ("BM-SS304-TUBE-4", "stainless steel", "304", "ornamental tube", "#4", 4.20, 4.75, 5.40),
    ("BM-SS304-BAR", "stainless steel", "304", "flat bar", "hot rolled annealed", 3.10, 3.50, 4.00),
    ("BM-AL3003-SH", "aluminum", "3003", "sheet", "mill", 3.05, 3.40, 3.90),
    ("BM-AL5052-SH", "aluminum", "5052", "sheet", "mill", 3.35, 3.75, 4.30),
    ("BM-AL5005-AN", "aluminum", "5005", "sheet", "clear anodized", 4.40, 4.95, 5.70),
    ("BM-AL6061-PL", "aluminum", "6061", "plate", "mill", 3.55, 3.95, 4.55),
    ("BM-AL6063-EXT", "aluminum", "6063", "extrusion", "mill", 3.40, 3.85, 4.40),
    ("BM-AL6061-EXT", "aluminum", "6061", "extrusion", "mill", 3.70, 4.15, 4.75),
    ("BM-AL3003-TREAD", "aluminum", "3003", "tread plate", "mill", 3.45, 3.85, 4.40),
    ("BM-GALV-G90", "galvanized steel", "CS Type B", "sheet", "G90", 0.88, 1.02, 1.20),
    ("BM-GALV-G60", "galvanized steel", "CS Type B", "sheet", "G60", 0.84, 0.97, 1.14),
    ("BM-CS-A36-PL", "carbon steel", "A36", "plate", "hot rolled", 0.78, 0.92, 1.08),
    ("BM-CS-A36-SHAPE", "carbon steel", "A36", "structural", "hot rolled", 0.82, 0.96, 1.12),
    ("BM-CS-A500-HSS", "carbon steel", "A500 Gr B/C", "HSS", "black", 0.90, 1.05, 1.24),
]
BENCH_MEDIAN = {b[0]: b[6] for b in BENCH}

MILLS = {
    ("stainless steel", "US"): "Riverbend Specialty Metals",
    ("stainless steel", "TW"): "Kaohsiung Harbor Stainless",
    ("stainless steel", "KR"): "Han River Steel Co.",
    ("stainless steel", "IN"): "Western Ghats Alloys",
    ("stainless steel", "CN"): "Jiangnan Huaxin Metals",
    ("aluminum", "US"): "Keystone Rolled Aluminum",
    ("aluminum", "CA"): "Lakehead Aluminum",
    ("galvanized steel", "US"): "Great Lakes Coated Steel",
    ("carbon steel", "US"): "Monongahela Structural",
}
DFARS_OK_MELT = {"US", "CA"}  # simplified for the demo

# --------------------------------------------------------------------------
# 2. Catalog
# --------------------------------------------------------------------------
CATALOG = []


def thou(x):
    return f"{round(x * 1000):03d}"


def num(x):
    return f"{x:g}"


def weight_each(shape, fam, t, w, h, L):
    d = DENSITY[fam]
    if shape in ("sheet", "plate", "flat bar"):
        vol = t * w * L
    elif shape == "tread plate":
        vol = t * w * L * 1.15  # raised pattern adds ~15%
    elif shape == "round tube":
        vol = math.pi * (w - t) * t * L
    elif shape in ("square tube", "rect tube", "HSS square"):
        vol = (2 * (w + h) - 4 * t) * t * L
    elif shape == "angle":
        vol = (w + h - t) * t * L
    elif shape == "channel":
        vol = (w + 2 * h - 2 * t) * t * L
    else:
        raise ValueError(shape)
    return round(vol * d, 2)


def add(code, fam, grade, temper, form, shape, finish, standard, dual_cert,
        t, w, h, L, bm, origin="US", melt=None, gauge="", certs="",
        factor=None, stock=None, lead=None, mtr="Y", recycled="auto",
        notes=""):
    melt = melt or origin
    if shape in ("sheet", "plate", "tread plate"):
        tag = f"{gauge}G" if gauge else thou(t)
        dim = f"{tag}-{num(w)}{num(L)}"
    elif shape == "round tube":
        dim = f"{round(w * 1000)}x{thou(t)}-{num(L)}"
    elif shape == "flat bar":
        dim = f"{thou(t)}x{num(w)}-{num(L)}"
    else:
        dim = f"{num(w)}x{num(h)}x{thou(t)}-{num(L)}"
    sku = f"{code}-{dim}-{origin}"

    if factor is None:
        factor = random.uniform(0.93, 1.04) if origin == "US" else random.uniform(0.86, 0.95)
    ppl = round(BENCH_MEDIAN[bm] * factor, 2)
    wt = weight_each(shape, fam, t, w, h, L)
    unit = round(wt * ppl, 2)

    if stock is None:
        stock = random.choice([0, 6, 12, 25, 40, 60, 90, 140]) if random.random() < 0.15 \
            else random.randint(15, 160)
    if lead is None:
        lead = 2 if stock > 0 else random.choice([14, 21, 28])

    if recycled == "auto":
        rng = {"stainless steel": (60, 82), "aluminum": (35, 70),
               "galvanized steel": (25, 35), "carbon steel": (85, 97)}[fam]
        recycled = "" if origin in ("CN", "IN") else random.randint(*rng)

    gauge_str = str(gauge) if gauge else ""
    # human description
    fin = f", {finish}" if finish else ""
    tmp = f"-{temper}" if temper else ""
    if shape in ("sheet", "plate", "tread plate"):
        thk = f'{gauge} ga ({t:.3f}")' if gauge else f'{t:.3f}"'
        desc = f'{grade}{tmp} {fam.title()} {shape.title()}{fin}, {thk} x {num(w)}" x {num(L)}"'
    elif shape == "round tube":
        desc = f'{grade} {fam.title()} Round Tube{fin}, {w:.3f}" OD x {t:.3f}" wall x {num(L / 12)}\''
    elif shape == "flat bar":
        desc = f'{grade}{tmp} {fam.title()} Flat Bar{fin}, {t:.3f}" x {num(w)}" x {num(L / 12)}\''
    else:
        desc = f'{grade}{tmp} {fam.title()} {shape.title()}{fin}, {num(w)}" x {num(h)}" x {t:.3f}" x {num(L / 12)}\''
    desc += f", {standard}"

    CATALOG.append({
        "sku": sku, "description": desc, "material_family": fam, "grade": grade,
        "temper": temper, "form": form, "shape": shape, "finish": finish,
        "standard": standard, "dual_cert": dual_cert, "gauge": gauge_str,
        "thickness_in": t, "width_in": w, "height_in": h if h else "",
        "length_in": L, "uom": "EA", "weight_lb_each": wt, "price_per_lb": ppl,
        "unit_price": unit, "unit_price_qty10": round(unit * 0.95, 2),
        "unit_price_qty50": round(unit * 0.91, 2), "benchmark_id": bm,
        "country_of_origin": origin, "country_of_melt": melt,
        "mill": MILLS.get((fam, origin), "Unknown"),
        "mtr_available": mtr,
        "mtr_type": ("EN 10204 3.1" if fam == "stainless steel" else "Mill cert (ASTM)") if mtr == "Y" else "",
        "buy_america_compliant": "Y" if origin == "US" and melt == "US" else "N",
        "dfars_compliant": "Y" if melt in DFARS_OK_MELT else "N",
        "certifications": certs or standard,
        "recycled_content_pct": recycled,
        "stock_qty": stock, "lead_time_days": lead, "min_order_qty": 1,
        "notes": notes,
    })
    return sku


SS = "stainless steel"
AL = "aluminum"
GV = "galvanized steel"
CS = "carbon steel"
A240 = "ASTM A240"


def ss_gauge(g):
    return GAUGES[g][0]


# 304 2B sheet
for g in [24, 22, 20]:
    add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
        ss_gauge(g), 48, 0, 96, "BM-SS304-SH-2B", origin="TW", gauge=g,
        certs="ASTM A240 / ASME SA-240")
add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
    ss_gauge(18), 48, 0, 96, "BM-SS304-SH-2B", origin="CN", gauge=18,
    certs="ASTM A240 / ASME SA-240")
for g in [16, 14, 11]:
    add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
        ss_gauge(g), 48, 0, 96, "BM-SS304-SH-2B", gauge=g, certs="ASTM A240 / ASME SA-240")
add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
    ss_gauge(16), 48, 0, 96, "BM-SS304-SH-2B", origin="TW", gauge=16, factor=0.88,
    certs="ASTM A240 / ASME SA-240")
for g in [18, 16, 14]:
    add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
        ss_gauge(g), 48, 0, 120, "BM-SS304-SH-2B", gauge=g, certs="ASTM A240 / ASME SA-240")
for g in [16, 14]:
    add("SS304-2B", SS, "304", "", "sheet", "sheet", "2B", A240, "304/304L",
        ss_gauge(g), 60, 0, 120, "BM-SS304-SH-2B", gauge=g, certs="ASTM A240 / ASME SA-240")

# 304 #4 sheet (PVC film one side)
for g in [20, 16, 14]:
    add("SS304-4", SS, "304", "", "sheet", "sheet", "#4", A240, "304/304L",
        ss_gauge(g), 48, 0, 96, "BM-SS304-SH-4", gauge=g,
        certs="ASTM A240 / ASME SA-240", notes="#4 brushed, PVC film one side")
add("SS304-4", SS, "304", "", "sheet", "sheet", "#4", A240, "304/304L",
    ss_gauge(18), 48, 0, 96, "BM-SS304-SH-4", origin="US", melt="MX", gauge=18,
    certs="ASTM A240 / ASME SA-240",
    notes="#4 brushed, PVC film. Polished in US from Mexican-melted coil")
add("SS304-4", SS, "304", "", "sheet", "sheet", "#4", A240, "304/304L",
    ss_gauge(16), 48, 0, 96, "BM-SS304-SH-4", origin="KR", gauge=16, factor=0.89,
    certs="ASTM A240 / ASME SA-240", notes="#4 brushed, PVC film one side")
for g in [18, 16]:
    add("SS304-4", SS, "304", "", "sheet", "sheet", "#4", A240, "304/304L",
        ss_gauge(g), 48, 0, 120, "BM-SS304-SH-4", gauge=g,
        certs="ASTM A240 / ASME SA-240", notes="#4 brushed, PVC film one side")

# 304 BA
for g in [18, 16]:
    add("SS304-BA", SS, "304", "", "sheet", "sheet", "BA", A240, "304/304L",
        ss_gauge(g), 48, 0, 96, "BM-SS304-SH-BA", origin="KR", gauge=g,
        certs="ASTM A240 / ASME SA-240", notes="Bright annealed, PVC film")

# 316
add("SS316-2B", SS, "316", "", "sheet", "sheet", "2B", A240, "316/316L",
    ss_gauge(18), 48, 0, 96, "BM-SS316-SH-2B", origin="IN", gauge=18,
    certs="ASTM A240 / ASME SA-240")
for g in [16, 14]:
    add("SS316-2B", SS, "316", "", "sheet", "sheet", "2B", A240, "316/316L",
        ss_gauge(g), 48, 0, 96, "BM-SS316-SH-2B", gauge=g, certs="ASTM A240 / ASME SA-240")
add("SS316-4", SS, "316", "", "sheet", "sheet", "#4", A240, "316/316L",
    ss_gauge(16), 48, 0, 96, "BM-SS316-SH-4", gauge=16, certs="ASTM A240 / ASME SA-240")
add("SS316-4", SS, "316", "", "sheet", "sheet", "#4", A240, "316/316L",
    ss_gauge(14), 48, 0, 120, "BM-SS316-SH-4", origin="KR", gauge=14, stock=0, lead=21,
    certs="ASTM A240 / ASME SA-240")
add("SS316-4", SS, "316", "", "sheet", "sheet", "#4", A240, "316/316L",
    ss_gauge(11), 48, 0, 96, "BM-SS316-SH-4", gauge=11, factor=1.16,
    certs="ASTM A240 / ASME SA-240", notes="Priced above market: small mill run")

# 304 plate
for t in [0.1875, 0.25, 0.375, 0.5]:
    add("SS304-PL", SS, "304", "", "plate", "plate", "No.1 HRAP", A240, "304/304L",
        t, 48, 0, 96, "BM-SS304-PL", certs="ASTM A240 / ASME SA-240")

# 304 ornamental tube (A554), #4 180 grit
A554 = "ASTM A554"
add("SS304-RT", SS, "304", "", "ornamental tube", "round tube", "#4", A554, "",
    TUBE_GAUGE_WALL[16], 1.5, 0, 240, "BM-SS304-TUBE-4", notes="#4 180-grit polished OD")
add("SS304-RT", SS, "304", "", "ornamental tube", "round tube", "#4", A554, "",
    TUBE_GAUGE_WALL[16], 1.66, 0, 240, "BM-SS304-TUBE-4", origin="IN",
    notes="#4 180-grit polished OD. Matches 1-1/4 NPS OD")
add("SS304-RT", SS, "304", "", "ornamental tube", "round tube", "#4", A554, "",
    TUBE_GAUGE_WALL[16], 1.9, 0, 240, "BM-SS304-TUBE-4",
    notes="#4 180-grit polished OD. Matches 1-1/2 NPS OD")
add("SS304-RT", SS, "304", "", "ornamental tube", "round tube", "#4", A554, "",
    TUBE_GAUGE_WALL[16], 2.0, 0, 240, "BM-SS304-TUBE-4", notes="#4 180-grit polished OD")
add("SS304-ST", SS, "304", "", "ornamental tube", "square tube", "#4", A554, "",
    TUBE_GAUGE_WALL[16], 1, 1, 240, "BM-SS304-TUBE-4", notes="#4 180-grit polished")
add("SS304-ST", SS, "304", "", "ornamental tube", "square tube", "#4", A554, "",
    TUBE_GAUGE_WALL[14], 2, 2, 240, "BM-SS304-TUBE-4", notes="#4 180-grit polished")

# 304 flat bar
for t, w in [(0.25, 1.5), (0.25, 2), (0.375, 2)]:
    add("SS304-FB", SS, "304", "", "flat bar", "flat bar", "hot rolled annealed",
        "ASTM A276 / A484", "304/304L", t, w, 0, 144, "BM-SS304-BAR", origin="IN")

# Aluminum sheet
B209 = "ASTM B209"
for t in [0.032, 0.040, 0.063, 0.080, 0.125]:
    add("AL3003-SH", AL, "3003", "H14", "sheet", "sheet", "mill", B209, "", t, 48, 0, 96, "BM-AL3003-SH")
add("AL3003-SH", AL, "3003", "H14", "sheet", "sheet", "mill", B209, "", 0.063, 48, 0, 144, "BM-AL3003-SH")
for t in [0.040, 0.063, 0.080, 0.090, 0.125, 0.190]:
    add("AL5052-SH", AL, "5052", "H32", "sheet", "sheet", "mill", B209, "", t, 48, 0, 96, "BM-AL5052-SH",
        origin="CA" if t in (0.090, 0.190) else "US")
for t in [0.063, 0.080, 0.125]:
    add("AL5052-SH", AL, "5052", "H32", "sheet", "sheet", "mill", B209, "", t, 48, 0, 120, "BM-AL5052-SH")
add("AL5052-SH", AL, "5052", "H32", "sheet", "sheet", "mill", B209, "", 0.125, 60, 0, 120, "BM-AL5052-SH")
ANOD = "ASTM B209; AAMA 611-14 Class I; AA-M12C22A41"
for t in [0.040, 0.063]:
    add("AL5005-AN", AL, "5005", "H34", "sheet", "sheet", "clear anodized", B209, "", t, 48, 0, 96,
        "BM-AL5005-AN", certs=ANOD, notes="Anodizing quality, clear Class I")
add("AL5005-AN", AL, "5005", "H34", "sheet", "sheet", "clear anodized", B209, "", 0.080, 48, 0, 96,
    "BM-AL5005-AN", certs=ANOD, factor=1.18, notes="Anodizing quality, clear Class I. Priced above market")
add("AL5005-AN", AL, "5005", "H34", "sheet", "sheet", "clear anodized", B209, "", 0.063, 48, 0, 120,
    "BM-AL5005-AN", certs=ANOD, notes="Anodizing quality, clear Class I")

# Aluminum plate
for t in [0.25, 0.375, 0.5]:
    add("AL6061-PL", AL, "6061", "T6", "plate", "plate", "mill", B209, "", t, 48, 0, 96, "BM-AL6061-PL",
        certs="ASTM B209 / AMS-QQ-A-250/11")
add("AL6061-PL", AL, "6061", "T651", "plate", "plate", "mill", B209, "", 0.5, 48, 0, 144, "BM-AL6061-PL",
    certs="ASTM B209 / AMS-QQ-A-250/11", notes="T651 stress relieved")

# Aluminum extrusions
B221 = "ASTM B221"
for w in [1, 1.5, 2]:
    add("AL6063-L", AL, "6063", "T52", "extrusion", "angle", "mill", B221, "", 0.125, w, w, 240, "BM-AL6063-EXT")
for w in [1, 2]:
    add("AL6063-SQ", AL, "6063", "T52", "extrusion", "square tube", "mill", B221, "", 0.125, w, w, 240, "BM-AL6063-EXT")
add("AL6063-REC", AL, "6063", "T52", "extrusion", "rect tube", "mill", B221, "", 0.125, 1, 2, 240, "BM-AL6063-EXT")
add("AL6063-REC", AL, "6063", "T52", "extrusion", "rect tube", "mill", B221, "", 0.125, 2, 4, 240, "BM-AL6063-EXT")
add("AL6063-REC", AL, "6063", "T52", "extrusion", "rect tube", "mill", B221, "", 0.125, 2, 4, 288, "BM-AL6063-EXT")
add("AL6063-C", AL, "6063", "T52", "extrusion", "channel", "mill", B221, "", 0.125, 2, 1, 240, "BM-AL6063-EXT")
add("AL6061-L", AL, "6061", "T6", "extrusion", "angle", "mill", B221, "", 0.25, 2, 2, 240, "BM-AL6061-EXT")

# Tread plate (no MTR, commercial quality)
for L in [96, 120]:
    add("AL3003-TP", AL, "3003", "H22", "tread plate", "tread plate", "mill (bright)", B209, "",
        0.125, 48, 0, L, "BM-AL3003-TREAD", mtr="N", certs="Commercial quality, no MTR",
        notes="Diamond / 5-bar pattern")

# Galvanized
A653 = "ASTM A653 CS Type B"
gv = lambda g: GAUGES[g][1]
for g in [24, 22, 20, 18, 16, 14]:
    add("GV90-SH", GV, "CS Type B", "", "sheet", "sheet", "G90", A653, "", gv(g), 48, 0, 96, "BM-GALV-G90",
        gauge=g, certs="ASTM A653 CS Type B, G90 coating")
for g in [20, 18, 16]:
    add("GV90-SH", GV, "CS Type B", "", "sheet", "sheet", "G90", A653, "", gv(g), 48, 0, 120, "BM-GALV-G90",
        gauge=g, certs="ASTM A653 CS Type B, G90 coating")
for g, L in [(20, 96), (16, 96), (16, 120)]:
    add("GV60-SH", GV, "CS Type B", "", "sheet", "sheet", "G60", A653, "", gv(g), 48, 0, L, "BM-GALV-G60",
        gauge=g, certs="ASTM A653 CS Type B, G60 coating")

# Carbon steel
for t in [0.25, 0.375, 0.5]:
    add("CSA36-PL", CS, "A36", "", "plate", "plate", "hot rolled", "ASTM A36", "", t, 48, 0, 96, "BM-CS-A36-PL",
        certs="ASTM A36 / A6")
for w in [2, 3]:
    add("CSA36-L", CS, "A36", "", "structural", "angle", "hot rolled", "ASTM A36", "", 0.25, w, w, 240,
        "BM-CS-A36-SHAPE", certs="ASTM A36 / A6")
for w in [2, 3]:
    add("CSA36-FB", CS, "A36", "", "structural", "flat bar", "hot rolled", "ASTM A36", "", 0.25, w, 0, 240,
        "BM-CS-A36-SHAPE", certs="ASTM A36 / A6")
add("HSS500-SQ", CS, "A500 Gr B/C", "", "HSS", "HSS square", "black", "ASTM A500", "", 0.1875, 2, 2, 240,
    "BM-CS-A500-HSS", certs="ASTM A500 Gr B/C dual")
add("HSS500-SQ", CS, "A500 Gr B/C", "", "HSS", "HSS square", "black", "ASTM A500", "", 0.1875, 2, 2, 288,
    "BM-CS-A500-HSS", certs="ASTM A500 Gr B/C dual", stock=6,
    notes="Only 6 in stock; balance 10 days")
add("HSS500-SQ", CS, "A500 Gr B/C", "", "HSS", "HSS square", "black", "ASTM A500", "", 0.1875, 3, 3, 240,
    "BM-CS-A500-HSS", certs="ASTM A500 Gr B/C dual")

SKUS = {r["sku"] for r in CATALOG}

# --------------------------------------------------------------------------
# 3. Synonyms, compliance rules, market index
# --------------------------------------------------------------------------
SYNONYMS = [
    # term, canonical, field, note
    ("SS", "stainless steel", "material", ""), ("S/S", "stainless steel", "material", ""),
    ("stainless", "stainless steel", "material", ""), ("SST", "stainless steel", "material", ""),
    ("T304", "304", "grade", ""), ("type 304", "304", "grade", ""), ("304L", "304", "grade", "stocked dual-cert 304/304L"),
    ("T316", "316", "grade", ""), ("316L", "316", "grade", "stocked dual-cert 316/316L"),
    ("2205", "duplex 2205", "grade", "NOT STOCKED"), ("duplex", "duplex 2205", "grade", "NOT STOCKED"),
    ("alum", "aluminum", "material", ""), ("alum.", "aluminum", "material", ""), ("AL", "aluminum", "material", ""),
    ("aluminium", "aluminum", "material", "UK spelling"),
    ("galv", "galvanized steel", "material", ""), ("galvanized", "galvanized steel", "material", ""),
    ("HDG", "galvanized steel", "material", "hot-dip galvanized; coating weight still needed"),
    ("CS", "carbon steel", "material", ""), ("HR", "carbon steel", "material", "hot rolled; default A36 for plate/bar/shapes"),
    ("MS", "carbon steel", "material", "mild steel; default A36"), ("black", "carbon steel", "material", "uncoated HSS/pipe"),
    ("brass", "brass", "material", "NOT STOCKED"),
    ("#4", "#4", "finish", "brushed / satin"), ("No. 4", "#4", "finish", ""), ("brushed", "#4", "finish", ""),
    ("satin", "#4", "finish", ""), ("polished", "#4", "finish", "ambiguous with #8 mirror; #8 not stocked"),
    ("mirror", "#8", "finish", "NOT STOCKED"), ("#8", "#8", "finish", "NOT STOCKED"),
    ("2B", "2B", "finish", "cold rolled mill finish"), ("mill finish", "2B (stainless) / mill (aluminum)", "finish", "depends on material"),
    ("BA", "BA", "finish", "bright annealed"), ("bright annealed", "BA", "finish", ""),
    ("clear anod", "clear anodized", "finish", ""), ("anod", "anodized", "finish", ""),
    ("diamond plate", "tread plate", "form", ""), ("checker plate", "tread plate", "form", ""),
    ("treadplate", "tread plate", "form", ""),
    ("SHT", "sheet", "form", ""), ("sht", "sheet", "form", ""), ("PL", "plate", "form", ""), ("PLT", "plate", "form", ""),
    ("L", "angle", "form", "L3x3x1/4 = angle 3 x 3 x 1/4"), ("C", "channel", "form", ""),
    ("HSS", "HSS square/rect tube", "form", "ASTM A500"), ("FB", "flat bar", "form", ""),
    ("4x8", "48 x 96", "size", "feet -> inches"), ("4' x 8'", "48 x 96", "size", ""), ("4x10", "48 x 120", "size", ""),
    ("5x10", "60 x 120", "size", ""), ("4x12", "48 x 144", "size", ""),
    ("1220x2440", "48 x 96", "size", "metric nominal, mm"), ("1250x2500", "49.2 x 98.4", "size", "metric, not stocked"),
    ("sticks", "lengths", "uom", ""), ("lgths", "lengths", "uom", ""), ("pcs", "EA", "uom", ""), ("sh", "EA (sheet)", "uom", ""),
    ("ga", "gauge", "thickness", "see gauge_table.csv; meaning differs by material"),
    ("1/8", "0.125 in", "thickness", "stainless 1/8 sheet is commonly 11 ga (0.120)"),
    ("MTR", "mill test report", "compliance", ""), ("mill cert", "mill test report", "compliance", ""),
    ("3.1 cert", "EN 10204 3.1", "compliance", ""), ("COO", "country of origin", "compliance", ""),
    ("Buy America", "Buy America (domestic melt + manufacture)", "compliance", ""),
    ("BABA", "Build America, Buy America", "compliance", ""), ("DFARS", "DFARS specialty metals", "compliance", ""),
    ("LEED", "recycled content documentation", "compliance", ""),
]

COMPLIANCE_RULES = [
    # rule_id, trigger, requirement, catalog_field, pass_condition, on_fail
    ("CR-01", "MTR / mill cert / test report", "Mill test report must ship with material",
     "mtr_available", "== 'Y'", "WARN: no MTR available for this SKU; ask buyer to accept C of C"),
    ("CR-02", "EN 10204 3.1", "Inspection certificate 3.1", "mtr_type", "== 'EN 10204 3.1'", "WARN"),
    ("CR-03", "Buy America / FTA / 49 CFR 661 / BABA", "Domestic melt AND manufacture",
     "buy_america_compliant", "== 'Y' (origin US and melt US)", "BLOCK: cannot quote as compliant; offer domestic alternative or ask"),
    ("CR-04", "DFARS", "Specialty metals melted in US or qualifying country", "dfars_compliant", "== 'Y'", "BLOCK"),
    ("CR-05", "COO / country of origin", "State country of origin per line", "country_of_origin", "not blank", "WARN"),
    ("CR-06", "Domestic preferred", "Prefer US origin when available", "country_of_origin", "== 'US'", "WARN: no domestic option"),
    ("CR-07", "AAMA 611", "Anodized finish certified to AAMA 611 (class as stated)", "certifications", "contains 'AAMA 611'", "BLOCK"),
    ("CR-08", "LEED / recycled content", "Recycled content % documented", "recycled_content_pct", "not blank", "WARN: documentation not available"),
    ("CR-09", "ASTM spec cited in RFQ", "Catalog standard must match the cited spec", "standard", "matches", "CLARIFY"),
    ("CR-10", "Required-by date", "Stock or lead time must meet date", "stock_qty / lead_time_days",
     "stock >= qty or lead fits", "WARN: partial availability"),
]

MARKET = [
    ("LME-NI", "Nickel (LME cash, SYNTHETIC)", "$/t", [15200, 15550, 15900, 16300, 16650, 16900]),
    ("LME-AL", "Aluminum (LME cash, SYNTHETIC)", "$/t", [2480, 2510, 2465, 2530, 2590, 2620]),
    ("MW-PREM", "US Midwest aluminum premium (SYNTHETIC)", "cents/lb", [62, 64, 63, 66, 68, 69]),
    ("US-HRC", "US hot rolled coil (SYNTHETIC)", "$/short ton", [880, 865, 850, 840, 855, 870]),
    ("SUR-304", "304 alloy surcharge (SYNTHETIC)", "$/lb", [0.78, 0.80, 0.83, 0.86, 0.89, 0.91]),
    ("SUR-316", "316 alloy surcharge (SYNTHETIC)", "$/lb", [1.38, 1.41, 1.45, 1.50, 1.54, 1.57]),
]
MONTHS = ["2026-05", "2026-06", "2026-07", "2026-08", "2026-09", "2026-10"]

# --------------------------------------------------------------------------
# 4. Sample RFQs + answer key
# --------------------------------------------------------------------------
KEY = []


def k(rfq, line, raw, status, skus, spec, qty, uom, compliance, question="", notes=""):
    for s in skus:
        assert s in SKUS, f"answer key SKU missing from catalog: {s}"
    KEY.append({
        "rfq_id": rfq, "line_no": line, "raw_text": raw, "expected_status": status,
        "expected_skus": ";".join(skus), "normalized_spec": spec, "qty": qty, "uom": uom,
        "compliance_result": compliance, "expected_clarifying_question": question, "notes": notes,
    })


# ---- RFQ-01: hotel lobby, casual email, private job
RFQ1_LINES = [
    "304 SS sheet 16ga 4x8, #4 finish w/ PVC, qty 24",
    "Stainless 304, 0.060 in, 48 x 96 -- 10 pcs",
    "1-1/2\" OD SS tube, 16 ga wall, #4 polished, 20' sticks x 12",
    "alum sheet .063 5052 4x10 qty 15",
    "1/4\" 6061 plate 4x8, need 4",
    "Galv sheet 18 ga 4x8 G90 - 30 sheets",
    "2x2 alum angle 1/8 wall, 20ft lengths, 40 pcs",
    "also need SS sheet 4x8 for the elevator cab panels, 8 sheets (architect still finalizing, use your best judgment)",
]
k("RFQ-01", 1, RFQ1_LINES[0], "MATCH", ["SS304-4-16G-4896-US", "SS304-4-16G-4896-KR"],
  "stainless 304 | sheet | #4 | 16ga 0.060 | 48x96 | ASTM A240", 24, "EA",
  "OK: MTR required, both SKUs have 3.1 MTR. No Buy America, so KR is eligible",
  notes="Two eligible SKUs; agent should pick and explain (price vs. stock vs. origin)")
k("RFQ-01", 2, RFQ1_LINES[1], "CLARIFY",
  ["SS304-2B-16G-4896-US", "SS304-2B-16G-4896-TW", "SS304-4-16G-4896-US", "SS304-4-16G-4896-KR", "SS304-BA-16G-4896-KR"],
  "stainless 304 | sheet | finish ? | 0.060 | 48x96", 10, "EA", "OK",
  "Which finish do you need for the 10 sheets of 304 0.060 x 48 x 96: 2B mill, #4 brushed (like line 1), or BA?",
  "Finish drives price ~12%; do not copy #4 from line 1 without asking")
k("RFQ-01", 3, RFQ1_LINES[2], "MATCH", ["SS304-RT-1500x065-240-US"],
  "stainless 304 | round tube | #4 | 1.500 OD x 0.065 wall | 240 in | ASTM A554", 12, "EA (20 ft length)",
  "OK: MTR available", notes="Tube 16 ga wall = 0.065, not sheet 16 ga 0.060")
k("RFQ-01", 4, RFQ1_LINES[3], "MATCH_ASSUMPTION", ["AL5052-SH-063-48120-US"],
  "aluminum 5052-H32 | sheet | mill | 0.063 | 48x120 | ASTM B209", 15, "EA", "OK",
  notes="Temper not stated; 5052 sheet only stocked as H32. Log assumption on quote")
k("RFQ-01", 5, RFQ1_LINES[4], "MATCH_ASSUMPTION", ["AL6061-PL-250-4896-US"],
  "aluminum 6061-T6 | plate | 0.250 | 48x96 | ASTM B209", 4, "EA", "OK",
  notes="Temper not stated; 6061 0.25 only stocked as T6")
k("RFQ-01", 6, RFQ1_LINES[5], "MATCH", ["GV90-SH-18G-4896-US"],
  "galvanized CS Type B | sheet | G90 | 18ga 0.0516 | 48x96 | ASTM A653", 30, "EA", "OK",
  notes="Galvanized gauge table, not stainless")
k("RFQ-01", 7, RFQ1_LINES[6], "MATCH_ASSUMPTION", ["AL6063-L-2x2x125-240-US"],
  "aluminum 6063-T52 | angle | 2 x 2 x 0.125 | 240 in | ASTM B221", 40, "EA (20 ft length)", "OK",
  notes="Alloy not stated; 2x2x1/8 angle only stocked in 6063-T52")
k("RFQ-01", 8, RFQ1_LINES[7], "CLARIFY", [],
  "stainless ? | sheet | finish ? | gauge ? | 48x96", 8, "EA", "OK: MTR required",
  "For the elevator cab panels: what grade (304 or 316), gauge and finish (#4, BA, or other)? "
  "We'll hold this line until the architect confirms.",
  "Buyer invites a guess. Agent must still ask; this is the trust moment")

# ---- RFQ-02: transit job, Buy America applies
RFQ2_LINES = [
    "(12) 316L #4 14ga 48x120",
    "(6) 3/8 x 4' x 8' A36 PL",
    "(10) L3x3x1/4 x 20'",
    "(8) HSS 2x2x3/16",
    "(4) 2205 duplex sht 10ga 48x96",
    "(20) 304 SS sheet 16ga 4x8 2B",
    "(20) 304 #4 18ga 4x8",
    "(25) galv 16ga 4x10",
]
k("RFQ-02", 1, RFQ2_LINES[0], "MATCH", ["SS316-4-14G-48120-KR"],
  "stainless 316/316L | sheet | #4 | 14ga 0.075 | 48x120 | ASTM A240", 12, "EA",
  "BLOCK: origin KR fails Buy America. Also 0 in stock, 21-day lead",
  notes="316L satisfied by dual cert. Must not quote as compliant")
k("RFQ-02", 2, RFQ2_LINES[1], "MATCH", ["CSA36-PL-375-4896-US"],
  "carbon steel A36 | plate | 0.375 | 48x96 | ASTM A36", 6, "EA", "OK: US melt + manufacture")
k("RFQ-02", 3, RFQ2_LINES[2], "MATCH", ["CSA36-L-3x3x250-240-US"],
  "carbon steel A36 | angle | 3 x 3 x 0.25 | 240 in", 10, "EA (20 ft length)", "OK")
k("RFQ-02", 4, RFQ2_LINES[3], "CLARIFY", ["HSS500-SQ-2x2x188-240-US", "HSS500-SQ-2x2x188-288-US"],
  "carbon steel A500 | HSS square | 2 x 2 x 0.1875 | length ?", 8, "EA", "OK",
  "What length do you need for the HSS 2x2x3/16: 20 ft or 24 ft? (Note: only 6 of the 24 ft are in stock.)")
k("RFQ-02", 5, RFQ2_LINES[4], "NO_MATCH", [],
  "duplex 2205 | sheet | 10ga | 48x96", 4, "EA", "n/a",
  "We don't stock duplex 2205. Should we source it (lead time TBD) or remove the line?",
  "Not in catalog; must not substitute 316")
k("RFQ-02", 6, RFQ2_LINES[5], "MATCH", ["SS304-2B-16G-4896-US"],
  "stainless 304 | sheet | 2B | 16ga 0.060 | 48x96 | ASTM A240", 20, "EA",
  "OK: must choose US SKU; TW SKU is cheaper but fails Buy America",
  notes="Compliance-driven SKU selection")
k("RFQ-02", 7, RFQ2_LINES[6], "MATCH", ["SS304-4-18G-4896-US"],
  "stainless 304 | sheet | #4 | 18ga 0.048 | 48x96", 20, "EA",
  "BLOCK: origin US but melt MX. Fails 'US melted' requirement",
  notes="Subtle trap: country_of_origin says US")
k("RFQ-02", 8, RFQ2_LINES[7], "CLARIFY", ["GV60-SH-16G-48120-US", "GV90-SH-16G-48120-US"],
  "galvanized | sheet | coating ? | 16ga | 48x120", 25, "EA", "OK",
  "Which zinc coating for the 16 ga galv 4x10: G60 or G90?")

# ---- RFQ-03: GC formal RFQ (PDF), residential amenity floor
RFQ3_LINES = [
    ("SS 304 SHT 1.5MM 1220X2440 BA", 6, "SH", "05 75 00 Decorative Formed Metal"),
    ("ALUM 5005-H34 CLR ANOD .063 4X8 (AAMA 611 CL I)", 20, "SH", "05 75 00 Decorative Formed Metal"),
    ("6063-T52 RECT TUBE 2 X 4 X .125 MILL 24'-0\"", 12, "EA", "05 50 00 Metal Fabrications"),
    ("6061-T651 PLT .500 X 48 X 144", 2, "EA", "05 50 00 Metal Fabrications"),
    ("SS FLAT BAR 1/4 X 2 TYPE 316 12'", 15, "EA", "05 73 00 Decorative Metal Railings"),
    ("BRASS SHT C260 1/2 HD .040 36X96", 8, "SH", "05 75 00 Decorative Formed Metal"),
    ("GALV SHT 20GA", 40, "SH", "07 62 00 Sheet Metal Flashing and Trim"),
    ("ALUM SHT .040 5052 MILL 4X8", 30, "SH", "07 62 00 Sheet Metal Flashing and Trim"),
]
k("RFQ-03", 1, RFQ3_LINES[0][0], "MATCH_ASSUMPTION", ["SS304-BA-16G-4896-KR"],
  "stainless 304 | sheet | BA | 1.5 mm = 0.059 in -> 16ga 0.060 | 1220x2440 mm = 48x96", 6, "EA",
  "WARN: state COO (KR) per GC requirement",
  notes="Metric to imperial: 1.6% thickness difference, within nominal; log conversion")
k("RFQ-03", 2, RFQ3_LINES[1][0], "MATCH", ["AL5005-AN-063-4896-US"],
  "aluminum 5005-H34 | sheet | clear anodized | 0.063 | 48x96", 20, "EA", "OK: AAMA 611 Class I certified")
k("RFQ-03", 3, RFQ3_LINES[2][0], "MATCH", ["AL6063-REC-2x4x125-288-US"],
  "aluminum 6063-T52 | rect tube | 2 x 4 x 0.125 | 288 in", 12, "EA (24 ft length)", "OK")
k("RFQ-03", 4, RFQ3_LINES[3][0], "MATCH", ["AL6061-PL-500-48144-US"],
  "aluminum 6061-T651 | plate | 0.500 | 48x144", 2, "EA", "OK")
k("RFQ-03", 5, RFQ3_LINES[4][0], "NO_MATCH", [],
  "stainless 316 | flat bar | 0.25 x 2 | 144 in", 15, "EA", "n/a",
  "We stock 1/4 x 2 flat bar in 304, not 316. Is 304 acceptable for this railing, or should we source 316?",
  "Alternative exists (SS304-FB-250x2-144-IN) but grade change needs buyer approval")
k("RFQ-03", 6, RFQ3_LINES[5][0], "NO_MATCH", [],
  "brass C260 | sheet | half hard | 0.040 | 36x96", 8, "EA", "n/a",
  "We don't carry brass. Please source this line elsewhere; we'll quote the rest.")
k("RFQ-03", 7, RFQ3_LINES[6][0], "CLARIFY", ["GV60-SH-20G-4896-US", "GV90-SH-20G-4896-US", "GV90-SH-20G-48120-US"],
  "galvanized | sheet | coating ? | 20ga | size ?", 40, "EA", "OK",
  "For the 20 ga galvanized: what sheet size (4x8 or 4x10) and coating (G60 or G90)?")
k("RFQ-03", 8, RFQ3_LINES[7][0], "MATCH_ASSUMPTION", ["AL5052-SH-040-4896-US"],
  "aluminum 5052-H32 | sheet | mill | 0.040 | 48x96", 30, "EA", "OK: recycled content documented",
  notes="Temper assumed H32")

# ---- RFQ-04: fabricator BOM spreadsheet
RFQ4_ROWS = [
    ("SS304-16-4X8-#4", 18, "sht", "stair stringer cladding"),
    ('Stainless steel sheet type 304 brushed, 0.0598" x 48" x 96"', 6, "ea", "added rev C"),
    ("Alum. diamond plate 3003 .125 4x8", 10, "sheets", "treads"),
    ("Alum sheet 0.04", "TBD", "", "guard infill?? confirm w/ architect"),
    ("SS tube 1.9 OD x .065 304 #4 20'", 10, "lengths", "handrail"),
    ("HR flat bar 1/4x2", 300, "ft", "pickets"),
    ("Stainless sheet 18ga 4x8 304 2B", 12, "ea", ""),
    ("SS 316 sheet 1/8 4x8 #4", 4, "ea", "exterior landing"),
]
k("RFQ-04", 1, RFQ4_ROWS[0][0], "MATCH", ["SS304-4-16G-4896-US", "SS304-4-16G-4896-KR"],
  "stainless 304 | sheet | #4 | 16ga | 48x96", 18, "EA", "WARN: domestic preferred -> pick US SKU if price gap is small")
k("RFQ-04", 2, RFQ4_ROWS[1][0], "MATCH_ASSUMPTION", ["SS304-4-16G-4896-US", "SS304-4-16G-4896-KR"],
  "stainless 304 | sheet | #4 | 0.0598 -> 16ga | 48x96", 6, "EA", "same as line 1",
  "Line 2 looks like the same item as line 1 (304 #4 16 ga 4x8). Combine into 24 sheets?",
  "0.0598 is the carbon-steel 16 ga decimal; interpret as 16 ga. Flag likely duplicate")
k("RFQ-04", 3, RFQ4_ROWS[2][0], "MATCH", ["AL3003-TP-125-4896-US"],
  "aluminum 3003-H22 | tread plate | 0.125 | 48x96", 10, "EA", "WARN: no MTR available (commercial quality)")
k("RFQ-04", 4, RFQ4_ROWS[3][0], "CLARIFY", ["AL3003-SH-040-4896-US", "AL5052-SH-040-4896-US", "AL5005-AN-040-4896-US"],
  "aluminum ? | sheet | 0.040 | size ? | qty ?", "TBD", "", "n/a",
  "For the 0.040 aluminum infill: which alloy/finish (3003 mill, 5052 mill, or 5005 clear anodized), sheet size, and quantity?")
k("RFQ-04", 5, RFQ4_ROWS[4][0], "MATCH", ["SS304-RT-1900x065-240-US"],
  "stainless 304 | round tube | #4 | 1.900 OD x 0.065 | 240 in", 10, "EA (20 ft length)", "OK")
k("RFQ-04", 6, RFQ4_ROWS[5][0], "MATCH_ASSUMPTION", ["CSA36-FB-250x2-240-US"],
  "carbon steel A36 | flat bar | 0.25 x 2 | 240 in", 15, "EA (20 ft length)", "OK",
  notes="300 ft / 20 ft lengths = 15 EA. HR assumed A36")
k("RFQ-04", 7, RFQ4_ROWS[6][0], "MATCH", ["SS304-2B-18G-4896-CN"],
  "stainless 304 | sheet | 2B | 18ga 0.048 | 48x96", 12, "EA",
  "WARN: domestic preferred but only CN origin stocked; no recycled-content doc")
k("RFQ-04", 8, RFQ4_ROWS[7][0], "MATCH_ASSUMPTION", ["SS316-4-11G-4896-US"],
  "stainless 316 | sheet | #4 | 1/8 -> 11ga 0.120 | 48x96", 4, "EA", "OK",
  notes="1/8 stainless sheet = 11 ga by trade convention. Our price is ABOVE benchmark high; agent must say so")


# --------------------------------------------------------------------------
# 5. Writers
# --------------------------------------------------------------------------
def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = fields or list(rows[0].keys())
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def write_data():
    cat = ROOT / "catalog"
    write_csv(cat / "catalog.csv", CATALOG)
    (cat / "catalog.json").write_text(json.dumps(CATALOG, indent=2), encoding="utf-8")

    write_csv(cat / "synonyms.csv", [dict(zip(["term", "canonical", "field", "note"], s)) for s in SYNONYMS])
    write_csv(cat / "gauge_table.csv", [
        {"gauge": g, "stainless_in": v[0], "galvanized_in": v[1], "carbon_steel_msg_in": v[2],
         "aluminum_bs_in": v[3]} for g, v in GAUGES.items()])
    write_csv(cat / "tube_wall_gauge.csv", [{"gauge": g, "wall_in": v} for g, v in TUBE_GAUGE_WALL.items()])
    write_csv(cat / "compliance_rules.csv", [dict(zip(
        ["rule_id", "trigger_in_rfq", "requirement", "catalog_field", "pass_condition", "on_fail"], r))
        for r in COMPLIANCE_RULES])

    bm = ROOT / "benchmarks"
    write_csv(bm / "price_benchmarks.csv", [{
        "benchmark_id": b[0], "material_family": b[1], "grade": b[2], "form": b[3], "finish_class": b[4],
        "low_per_lb": b[5], "median_per_lb": b[6], "high_per_lb": b[7], "unit": "USD/lb",
        "basis": "FOB distributor, 1,000-5,000 lb order", "small_order_adder_pct": "8-15 (<500 lb)",
        "domestic_premium_pct": "6-12 vs import", "as_of": AS_OF, "source": SYNTH} for b in BENCH])
    rows = []
    for idx, name, unit, vals in MARKET:
        for m, v in zip(MONTHS, vals):
            rows.append({"index_id": idx, "name": name, "unit": unit, "month": m, "value": v, "source": SYNTH})
    write_csv(bm / "market_index.csv", rows)

    write_csv(ROOT / "rfqs" / "answer_key.csv", KEY)

    # SQLite: the "fake database"
    db = ROOT / "rudco_demo.db"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    for table, path in [("catalog", cat / "catalog.csv"), ("synonyms", cat / "synonyms.csv"),
                        ("gauge_table", cat / "gauge_table.csv"), ("tube_wall_gauge", cat / "tube_wall_gauge.csv"),
                        ("compliance_rules", cat / "compliance_rules.csv"),
                        ("price_benchmarks", bm / "price_benchmarks.csv"), ("market_index", bm / "market_index.csv"),
                        ("answer_key", ROOT / "rfqs" / "answer_key.csv")]:
        with open(path, encoding="utf-8") as f:
            r = list(csv.DictReader(f))
        cols = list(r[0].keys())
        coldefs = ", ".join('"' + c + '" TEXT' for c in cols)
        con.execute(f"CREATE TABLE {table} ({coldefs})")
        con.executemany(f"INSERT INTO {table} VALUES ({', '.join('?' * len(cols))})",
                        [tuple(x[c] for c in cols) for x in r])
    con.commit()
    con.close()


def write_emails():
    out = ROOT / "rfqs"
    m = EmailMessage()
    m["From"] = "Dana Whitfield <dana@bowerymillwork.example>"
    m["To"] = "quotes@gowanusmetal.example"
    m["Date"] = "Thu, 08 Oct 2026 16:42:10 -0400"
    m["Subject"] = "RFQ - Lafayette St hotel lobby pkg - need #s by Fri"
    body = ("Hi team,\n\n"
            "Can you price the below for the Lafayette St hotel lobby job? Mostly reception desk cladding "
            "+ elevator lobby trim. Delivery to our shop in Red Hook (Brooklyn 11231).\n\n"
            + "\n".join(f"- {l}" for l in RFQ1_LINES) +
            "\n\nMill certs required on all stainless. Private job, no Buy America. Need material by 11/2.\n\n"
            "Thanks,\nDana\n\n--\nDana Whitfield | Purchasing\nBowery Millwork & Fixtures\n")
    m.set_content(body, cte="8bit")
    (out / "rfq_01_hotel_lobby_email.eml").write_bytes(bytes(m))

    m = EmailMessage()
    m["From"] = "Marcus Oyelaran <moyelaran@halseyornamental.example>"
    m["To"] = "quotes@gowanusmetal.example"
    m["Date"] = "Fri, 09 Oct 2026 07:05:44 -0400"
    m["Subject"] = "FW: RFQ#HO-2291 transit canopy misc metals"
    body = ("pls quote ASAP. FTA funded -> Buy America applies (49 CFR 661), all steel/iron must be US "
            "melted+manufactured, stainless too per spec 05 50 00 2.1.B. MTRs + COO cert w/ each line. "
            "Quote valid 30 days min pls.\n\n"
            + "\n".join(RFQ2_LINES) +
            "\n\nthx\nM.\n\nSent from my iPhone\n\n"
            "---------- Forwarded message ---------\n"
            "From: Priya N. <pnair@stationrehab-gc.example>\n"
            "Subject: RFQ#HO-2291 transit canopy misc metals\n\n"
            "Marcus - see attached takeoff, need supplier #s by Wed. Buy America cert is a hard requirement, "
            "no waivers on this one.\n")
    m.set_content(body, cte="8bit")
    (out / "rfq_02_transit_buy_america_email.eml").write_bytes(bytes(m))


def write_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path = ROOT / "rfqs" / "rfq_03_residential_gc_form.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title="RFQ AAB-0417", author="Atlantic Avenue Builders LLC (fictional)")
    ss = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=8.5, leading=11)
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=8.5, leading=10.5)
    story = [
        Paragraph("<b>ATLANTIC AVENUE BUILDERS LLC</b>", ParagraphStyle("h", parent=ss["Title"], fontSize=15,
                                                                        alignment=0, spaceAfter=2)),
        Paragraph("General Contractor &nbsp;|&nbsp; 455 Atlantic Ave, Brooklyn NY 11217 &nbsp;|&nbsp; "
                  "procurement@aabuilders.example", small),
        Spacer(1, 10),
        Paragraph("<b>REQUEST FOR QUOTATION - MATERIALS ONLY</b>", ss["Heading3"]),
    ]
    meta = [["RFQ No.", "AAB-0417", "Issued", "10/09/2026"],
            ["Project", "210 Bergen St Residential - Amenity Level", "Quotes due", "10/16/2026 5:00 PM"],
            ["Deliver to", "210 Bergen St, Brooklyn NY (site, no forklift)", "Buyer", "J. Castellano"]]
    t = Table(meta, colWidths=[0.9 * inch, 3.3 * inch, 0.9 * inch, 2.1 * inch])
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8.5), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                           ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(1, 12)]
    rows = [["Line", "Description", "Qty", "U/M", "Spec Section", "Unit Price", "Ext."]]
    for i, (d, q, u, sec) in enumerate(RFQ3_LINES, 1):
        rows.append([str(i), Paragraph(d, cell), str(q), u, Paragraph(sec, cell), "", ""])
    t = Table(rows, colWidths=[0.4 * inch, 2.9 * inch, 0.45 * inch, 0.45 * inch, 1.6 * inch, 0.75 * inch, 0.65 * inch],
              repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2b2b2b")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                           ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")])]))
    story += [t, Spacer(1, 14), Paragraph("<b>GENERAL REQUIREMENTS</b>", ss["Heading5"])]
    reqs = [
        "1. Provide mill test reports (MTRs) for ALL metal items with delivery.",
        "2. State country of origin for each line on the quotation.",
        "3. Project is pursuing LEED v4.1. Provide recycled content (pre-/post-consumer %) documentation per line.",
        "4. Anodized finishes to comply with AAMA 611, Class I clear.",
        "5. Quote firm for 30 days. Freight to be shown separately.",
        "6. Substitutions require written approval from Architect. Do not substitute alloys or grades.",
        "7. Dimensions per architectural drawings A-501 through A-508 (metric details on A-504).",
    ]
    story += [Paragraph(r, small) for r in reqs]
    story += [Spacer(1, 16), Paragraph("Supplier: ______________________ &nbsp;&nbsp; Signature: ______________ "
                                       "&nbsp;&nbsp; Date: __________", small)]
    doc.build(story)


def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM rev C"
    ws["A1"] = "NAVY YARD FABRICATION CO. - MATERIAL TAKEOFF / RFQ"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")
    ws["A2"] = "Job: 41-S Residential Stair + Guards     Rev C     10/07/26     by: T.R."
    ws.merge_cells("A2:F2")
    ws["A3"] = "Certs: MTRs where available. Domestic preferred, not required. Pls send quote to tr@navyyardfab.example"
    ws.merge_cells("A3:F3")
    ws["A3"].font = Font(italic=True, color="9C0006")
    hdr = ["Item", "Desc", "Qty", "Units", "Notes", "Price (supplier fill in)"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    thin = Side(style="thin", color="BFBFBF")
    for i, (d, q, u, n) in enumerate(RFQ4_ROWS, 1):
        r = 5 + i + (1 if i > 4 else 0)  # stray blank row after item 4
        vals = [i, d, q, u, n, None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 2:
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="FFF2CC")  # rev C highlight
        if i == 4:
            ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor="FFC7CE")
    ws.cell(row=16, column=1, value="*** line 8 per arch RFI #14 - may change ***").font = Font(italic=True, size=9)
    for col, w in zip("ABCDEF", [6, 52, 8, 9, 30, 22]):
        ws.column_dimensions[col].width = w

    old = wb.create_sheet("rev B (old)")
    old["A1"] = "SUPERSEDED - see rev C"
    old["A1"].font = Font(bold=True, color="FF0000")
    old.append([])
    old.append(["Item", "Desc", "Qty", "Units"])
    old.append([1, "SS304-16-4X8-#4", 12, "sht"])
    old.append([2, "Alum. diamond plate 3003 .125 4x8", 8, "sheets"])
    wb.save(ROOT / "rfqs" / "rfq_04_fabricator_bom.xlsx")


if __name__ == "__main__":
    (ROOT / "rfqs").mkdir(parents=True, exist_ok=True)
    write_data()
    write_emails()
    write_pdf()
    write_xlsx()
    print(f"catalog SKUs: {len(CATALOG)}  benchmarks: {len(BENCH)}  answer-key lines: {len(KEY)}")
